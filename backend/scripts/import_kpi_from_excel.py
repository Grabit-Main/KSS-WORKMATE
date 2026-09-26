import os
import sys
from datetime import datetime
import openpyxl

backend_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

from app.database import Base, engine, SessionLocal
from app.models.kpi import DailyKPILog
from app.models.user import User

def normalize_name(name: str) -> str:
    return "".join(c.lower() for c in str(name) if c.isalnum())

def seed_kpi():
    print("1. Verifying database tables...")
    Base.metadata.create_all(bind=engine)

    db = SessionLocal()
    try:
        tl = db.query(User).filter(User.role == "TL").first()
        if not tl:
            tl = db.query(User).filter(User.email.ilike("%satya%")).first()
        if not tl:
            tl = db.query(User).first()
        print(f"Using Evaluator (TL): {tl.full_name} ({tl.role}, {tl.id})")

        users = db.query(User).all()
        user_map = {}
        for u in users:
            norm_full = normalize_name(u.full_name)
            user_map[norm_full] = u
            user_map[normalize_name(f"{u.first_name}{u.last_name}")] = u
            user_map[normalize_name(u.first_name)] = u

        excel_paths = [
            r"C:\Users\dassa\Downloads\Employee_KPI_Tracker- Final Sheet.xlsx",
            r"C:\Users\dassa\Downloads\Employee_KPI_Tracker-01.xlsx"
        ]

        target_path = None
        for p in excel_paths:
            if os.path.exists(p):
                target_path = p
                break

        if not target_path:
            print("No KPI excel file found.")
            return

        print(f"Loading KPI data from: {target_path}")
        wb = openpyxl.load_workbook(target_path, data_only=True)
        if "Daily KPI Log" not in wb.sheetnames:
            print("Sheet 'Daily KPI Log' not found.")
            return

        sheet = wb["Daily KPI Log"]
        seeded = 0
        skipped = 0

        for row in list(sheet.iter_rows(values_only=True))[3:]:
            if not row or len(row) < 13:
                continue
            date_val = row[0]
            emp_name = row[1]
            if not date_val or not emp_name:
                continue

            if isinstance(date_val, datetime):
                log_date = date_val.date()
            else:
                date_str = str(date_val).strip()
                try:
                    log_date = datetime.strptime(date_str, "%Y-%m-%d %H:%M:%S").date()
                except ValueError:
                    try:
                        log_date = datetime.strptime(date_str, "%Y-%m-%d").date()
                    except ValueError:
                        continue

            norm_emp = normalize_name(emp_name)
            matched_user = user_map.get(norm_emp)

            if not matched_user:
                for k, u in user_map.items():
                    if norm_emp in k or k in norm_emp:
                        matched_user = u
                        break

            if not matched_user:
                skipped += 1
                continue

            try:
                scores = [int(float(x)) for x in row[2:12]]
            except Exception:
                continue

            pct, month_str, status_str = DailyKPILog.calculate_metrics(
                scores[0], scores[1], scores[2], scores[3], scores[4],
                scores[5], scores[6], scores[7], scores[8], scores[9],
                log_date
            )

            existing = db.query(DailyKPILog).filter(
                DailyKPILog.employee_id == matched_user.id,
                DailyKPILog.date == log_date
            ).first()

            if not existing:
                kpi_log = DailyKPILog(
                    date=log_date,
                    employee_id=matched_user.id,
                    evaluator_id=tl.id,
                    task_completion=scores[0],
                    quality=scores[1],
                    productivity=scores[2],
                    deadline_adherence=scores[3],
                    ownership=scores[4],
                    problem_solving=scores[5],
                    communication=scores[6],
                    team_collaboration=scores[7],
                    learning_improvement=scores[8],
                    attendance_discipline=scores[9],
                    daily_kpi_percentage=pct,
                    month=month_str,
                    status=status_str,
                    notes=f"Daily performance evaluation for {matched_user.full_name}"
                )
                db.add(kpi_log)
                seeded += 1

        db.commit()
        print(f"KPI Seeding Completed! Seeded {seeded} daily KPI records successfully into database.")

    except Exception as e:
        db.rollback()
        print(f"Error seeding KPI data: {e}")
        raise
    finally:
        db.close()

if __name__ == "__main__":
    seed_kpi()
